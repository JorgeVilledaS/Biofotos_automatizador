"""
Selector de fotos de especies forestales.

Uso:
    pip install streamlit requests pandas python-docx
    python -m streamlit run app.py

Si existe fotos.csv (el que genera el script anterior), se usa como fuente
de candidatas. Si una especie no está ahí, se consulta iNaturalist en vivo.
Las elecciones se guardan en elegidas.csv, así que se puede cerrar y retomar.
"""
import io
import os

import pandas as pd
import requests
import streamlit as st
from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.shared import Cm, Pt
from PIL import Image

CSV_FOTOS = "fotos.csv"
CSV_ELEGIDAS = "elegidas.csv"
COLS_ELEGIDAS = ["especie", "aceptado", "estado", "url", "licencia", "autor", "obs"]


# ---------------------------------------------------------------- datos ----
@st.cache_data(show_spinner=False)
def nombre_aceptado(nombre):
    try:
        r = requests.get(
            "https://api.gbif.org/v1/species/match",
            params={"name": nombre, "kingdom": "Plantae"},
            timeout=20,
        ).json()
        return r.get("species") or nombre
    except Exception:
        return nombre


@st.cache_data(show_spinner=False)
def fotos_inat(nombre, max_fotos, place_id, licencias):
    params = {
        "taxon_name": nombre,
        "quality_grade": "research",
        "photo_license": licencias,
        "per_page": max_fotos,
        "order_by": "votes",
    }
    if place_id:
        params["place_id"] = place_id
    try:
        r = requests.get(
            "https://api.inaturalist.org/v1/observations", params=params, timeout=30
        ).json()
    except Exception:
        return []
    out = []
    for obs in r.get("results", []):
        if not obs.get("photos"):
            continue
        p = obs["photos"][0]
        out.append(
            {
                "url": p["url"].replace("square", "large"),
                "licencia": p.get("license_code") or "",
                "autor": p.get("attribution") or "",
                "obs": obs.get("uri", ""),
            }
        )
    return out


@st.cache_data(show_spinner=False)
def cargar_fotos_csv(mtime):
    if not os.path.exists(CSV_FOTOS):
        return pd.DataFrame()
    df = pd.read_csv(CSV_FOTOS).fillna("")
    return df


def candidatas(especie, max_fotos, place_id, licencias, fallback_mundo):
    """Primero fotos.csv; si no hay, API en vivo."""
    mtime = os.path.getmtime(CSV_FOTOS) if os.path.exists(CSV_FOTOS) else 0
    df = cargar_fotos_csv(mtime)
    if not df.empty:
        e = especie.strip().lower()
        m = df[
            (df["original"].str.lower() == e) | (df["aceptado"].str.lower() == e)
        ]
        if not m.empty:
            return m[["url", "licencia", "autor", "obs"]].to_dict("records"), especie
    acept = nombre_aceptado(especie)
    fotos = fotos_inat(acept, max_fotos, place_id, licencias)
    if not fotos and fallback_mundo and place_id:
        fotos = fotos_inat(acept, max_fotos, 0, licencias)
    return fotos, acept


@st.cache_data(show_spinner=False)
def descargar_img(url):
    try:
        r = requests.get(url, timeout=30)
        if r.status_code == 200:
            return r.content
    except Exception:
        pass
    return None


# ------------------------------------------------------------ documento ----
def _poner_imagen(par, data, max_w_cm, max_h_cm):
    with Image.open(io.BytesIO(data)) as im:
        w, h = im.size
    esc = min(max_w_cm / w, max_h_cm / h)
    par.add_run().add_picture(io.BytesIO(data), width=Cm(w * esc), height=Cm(h * esc))


def _ficha(contenedor, fila, max_w, max_h, centrar=True):
    """Escribe imagen + nombre + crédito en un documento o celda."""
    par = contenedor.add_paragraph()
    if centrar:
        par.alignment = WD_ALIGN_PARAGRAPH.CENTER
    data = descargar_img(fila["url"])
    if data:
        _poner_imagen(par, data, max_w, max_h)
    else:
        par.add_run("(imagen no disponible)")
    p2 = contenedor.add_paragraph()
    if centrar:
        p2.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = p2.add_run(fila["aceptado"] or fila["especie"])
    r.italic = True
    r.bold = True
    r.font.size = Pt(12)
    p3 = contenedor.add_paragraph()
    if centrar:
        p3.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r3 = p3.add_run(f"Foto: {fila['autor']} ({fila['licencia'].upper()})")
    r3.font.size = Pt(8)


def construir_docx(filas, por_pagina, titulo):
    doc = Document()
    sec = doc.sections[0]
    for lado in ("left_margin", "right_margin", "top_margin", "bottom_margin"):
        setattr(sec, lado, Cm(1.8))
    if titulo:
        doc.add_heading(titulo, level=0)

    if por_pagina == 1:
        for i, fila in enumerate(filas):
            _ficha(doc, fila, 16, 20)
            if i < len(filas) - 1:
                doc.add_page_break()
    else:  # cuadrícula 2x2
        for inicio in range(0, len(filas), 4):
            grupo = filas[inicio : inicio + 4]
            tabla = doc.add_table(rows=2, cols=2)
            for k, fila in enumerate(grupo):
                celda = tabla.cell(k // 2, k % 2)
                celda.paragraphs[0]._element.getparent().remove(
                    celda.paragraphs[0]._element
                )
                _ficha(celda, fila, 8, 10)
            if inicio + 4 < len(filas):
                doc.add_page_break()
    buf = io.BytesIO()
    doc.save(buf)
    return buf.getvalue()


# ===== UI =====
st.set_page_config(page_title="Selector de fotos", layout="wide")


def cargar_elegidas():
    if os.path.exists(CSV_ELEGIDAS):
        df = pd.read_csv(CSV_ELEGIDAS).fillna("")
        return {r["especie"]: r for r in df.to_dict("records")}
    return {}


def guardar_elegidas():
    filas = list(st.session_state.elegidas.values())
    pd.DataFrame(filas, columns=COLS_ELEGIDAS).to_csv(
        CSV_ELEGIDAS, index=False, encoding="utf-8-sig"
    )


def elegir(especie, acept, foto):
    st.session_state.elegidas[especie] = {
        "especie": especie,
        "aceptado": acept,
        "estado": "elegida",
        "url": foto["url"],
        "licencia": foto["licencia"],
        "autor": foto["autor"],
        "obs": foto["obs"],
    }
    guardar_elegidas()


def ninguna(especie, acept):
    st.session_state.elegidas[especie] = {
        "especie": especie, "aceptado": acept, "estado": "ninguna",
        "url": "", "licencia": "", "autor": "", "obs": "",
    }
    guardar_elegidas()


def quitar(especie):
    st.session_state.elegidas.pop(especie, None)
    guardar_elegidas()


def ir(delta):
    st.session_state.idx = max(0, min(len(st.session_state.cola) - 1, st.session_state.idx + delta))


if "elegidas" not in st.session_state:
    st.session_state.elegidas = cargar_elegidas()
    # al reabrir, se recuperan las especies ya trabajadas
    st.session_state.cola = list(st.session_state.elegidas.keys())
    st.session_state.idx = 0

# ---- barra lateral
with st.sidebar:
    st.header("Especies")
    modo = st.radio("Cómo cargar", ["Una por una", "Archivo .txt", "Todas de fotos.csv"])
    if modo == "Una por una":
        nueva = st.text_input("Nombre científico", placeholder="Ceiba pentandra")
        if st.button("Buscar", type="primary") and nueva.strip():
            n = nueva.strip()
            if n not in st.session_state.cola:
                st.session_state.cola.append(n)
            st.session_state.idx = st.session_state.cola.index(n)
    elif modo == "Archivo .txt":
        arch = st.file_uploader("Una especie por línea", type="txt")
        if arch and st.button("Cargar lista", type="primary"):
            lineas = arch.read().decode("utf-8-sig").splitlines()
            nuevas = [l.strip() for l in lineas if l.strip()]
            for n in nuevas:
                if n not in st.session_state.cola:
                    st.session_state.cola.append(n)
            st.session_state.idx = st.session_state.cola.index(nuevas[0]) if nuevas else 0
    else:
        if st.button("Cargar", type="primary"):
            df = cargar_fotos_csv(os.path.getmtime(CSV_FOTOS) if os.path.exists(CSV_FOTOS) else 0)
            if df.empty:
                st.warning("No encontré fotos.csv en esta carpeta.")
            else:
                for n in df["original"].drop_duplicates():
                    if n not in st.session_state.cola:
                        st.session_state.cola.append(n)

    st.divider()
    st.subheader("Búsqueda en iNaturalist")
    st.caption("Solo se usa para especies que no estén en fotos.csv.")
    max_fotos = st.slider("Candidatas por especie", 3, 12, 6)
    place_id = st.number_input("place_id (0 = todo el mundo)", value=6940, step=1)
    fallback = st.checkbox("Si no hay fotos en ese lugar, buscar en todo el mundo", True)
    licencias = st.multiselect(
        "Licencias permitidas", ["cc0", "cc-by", "cc-by-sa"], ["cc0", "cc-by", "cc-by-sa"]
    )
    lic_str = ",".join(licencias) or "cc0,cc-by,cc-by-sa"

    st.divider()
    if st.button("Vaciar lista de trabajo"):
        st.session_state.cola = []
        st.session_state.idx = 0

# ---- pestañas
tab_elegir, tab_doc = st.tabs(["1. Elegir fotos", "2. Documento para imprimir"])

with tab_elegir:
    cola = st.session_state.cola
    if not cola:
        st.info("Agrega una especie o sube un .txt en la barra lateral.")
    else:
        hechas = sum(1 for e in cola if e in st.session_state.elegidas)
        st.progress(hechas / len(cola), text=f"{hechas} de {len(cola)} especies revisadas")

        especie = cola[st.session_state.idx]
        c1, c2, c3 = st.columns([1, 6, 1])
        c1.button("← Anterior", on_click=ir, args=(-1,), disabled=st.session_state.idx == 0)
        c3.button(
            "Siguiente →", on_click=ir, args=(1,),
            disabled=st.session_state.idx >= len(cola) - 1,
        )
        c2.subheader(f"{st.session_state.idx + 1}/{len(cola)} · *{especie}*")

        with st.spinner("Buscando fotos..."):
            fotos, acept = candidatas(especie, max_fotos, int(place_id), lic_str, fallback)

        if acept != especie:
            st.caption(f"Nombre aceptado actual (GBIF): *{acept}*")

        actual = st.session_state.elegidas.get(especie)
        if actual and actual["estado"] == "elegida":
            st.success("Ya elegiste una foto para esta especie (marcada abajo).")
        elif actual:
            st.warning("Marcada como «ninguna sirve».")

        if not fotos:
            st.error("No se encontraron fotos. Puedes marcarla para buscar a mano.")
        else:
            columnas = st.columns(3)
            for i, f in enumerate(fotos):
                with columnas[i % 3]:
                    st.image(f["url"], use_container_width=True)
                    marcada = bool(actual and actual["url"] == f["url"])
                    st.caption(f"{'✅ ' if marcada else ''}{f['autor']} · {f['licencia'].upper()}")
                    st.button(
                        "Elegida" if marcada else "Elegir esta",
                        key=f"el_{especie}_{i}",
                        on_click=elegir, args=(especie, acept, f),
                        type="primary" if marcada else "secondary",
                    )

        b1, b2, _ = st.columns([2, 2, 6])
        b1.button("Ninguna sirve", on_click=ninguna, args=(especie, acept))
        b2.button("Quitar elección", on_click=quitar, args=(especie,))

with tab_doc:
    elegidas = [
        st.session_state.elegidas[e]
        for e in st.session_state.cola
        if e in st.session_state.elegidas and st.session_state.elegidas[e]["estado"] == "elegida"
    ]
    pendientes = [e for e in st.session_state.cola if e not in st.session_state.elegidas]
    sin_foto = [
        e for e in st.session_state.cola
        if e in st.session_state.elegidas and st.session_state.elegidas[e]["estado"] == "ninguna"
    ]

    st.write(f"**{len(elegidas)}** fotos elegidas · **{len(pendientes)}** pendientes · **{len(sin_foto)}** sin foto adecuada")
    if sin_foto:
        st.caption("Buscar a mano: " + ", ".join(sin_foto))

    titulo = st.text_input("Título del documento (opcional)", "Especies forestales autóctonas de Guatemala")
    layout = st.radio("Fotos por página", ["1 grande por página", "4 por página (2×2)"], horizontal=True)

    if elegidas:
        if st.button("Generar documento", type="primary"):
            with st.spinner("Descargando imágenes y armando el documento..."):
                datos = construir_docx(elegidas, 1 if layout.startswith("1") else 4, titulo)
            st.session_state.docx = datos
        if "docx" in st.session_state:
            st.download_button(
                "Descargar .docx",
                st.session_state.docx,
                file_name="guia_especies.docx",
                mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            )
            st.caption("Ábrelo en Word o LibreOffice para imprimir o exportar a PDF.")
    else:
        st.info("Elige al menos una foto en la pestaña anterior.")
